"""Build the PowerShield-FL conference-paper DOCX from the Markdown source.

The Markdown file is the editable source.  This builder intentionally supports
only the subset of Markdown used in the paper: headings, paragraphs, pipe
tables, bullet/numbered lists, block quotes, images, display equations, inline
bold/italic/code, and reference entries.  That keeps the render deterministic
without adding a heavy Markdown dependency.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "PowerShield_FL_Conference_Paper.md"
OUTPUT = ROOT / "PowerShield_FL_Conference_Paper.docx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x5C, 0x8A)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x59, 0x59, 0x59)
LIGHT_GRAY = RGBColor(0x7F, 0x7F, 0x7F)
RED = RGBColor(0xB0, 0x3A, 0x2E)
GREEN = RGBColor(0x2E, 0x8B, 0x57)

BODY_FONT = "Cambria"
HEAD_FONT = "Cambria"
MONO_FONT = "Consolas"

FIGURE_WIDTHS = {
    "paper_architecture.png": 6.7,
    "paper_convergence.png": 6.4,
    "paper_communication.png": 6.4,
    "paper_privacy.png": 6.4,
    "paper_scalability.png": 6.4,
    "paper_feature_importance.png": 6.0,
    "paper_confusion.png": 6.0,
    "paper_window_vs_consumer.png": 6.2,
}


# --------------------------------------------------------------------------- #
# Low-level helpers
# --------------------------------------------------------------------------- #
def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=40, start=70, bottom=40, end=70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="C9D4E3", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def set_table_cell_margins(table, top=30, start=60, bottom=30, end=60) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_bottom_border(paragraph, color="1F3864", size="10") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_run(paragraph, text: str, *, bold=False, italic=False, size=None, color=None, font=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if font is not None:
        run.font.name = font
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return run


def add_inline(paragraph, text: str, *, size=10.4, color=DARK, bold=False, italic=False, font=BODY_FONT):
    """Render a small subset of inline Markdown: bold, italic, code."""
    token = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
    pos = 0
    for match in token.finditer(text):
        if match.start() > pos:
            add_run(paragraph, text[pos:match.start()], bold=bold, italic=italic, size=size, color=color, font=font)
        piece = match.group(0)
        if piece.startswith("**") and piece.endswith("**"):
            add_run(paragraph, piece[2:-2], bold=True, italic=italic, size=size, color=color, font=font)
        elif piece.startswith("`") and piece.endswith("`"):
            add_run(paragraph, piece[1:-1], bold=bold, italic=False, size=size - 0.6, color=color, font=MONO_FONT)
        elif piece.startswith("*") and piece.endswith("*"):
            add_run(paragraph, piece[1:-1], bold=bold, italic=True, size=size, color=color, font=font)
        pos = match.end()
    if pos < len(text):
        add_run(paragraph, text[pos:], bold=bold, italic=italic, size=size, color=color, font=font)


def add_hyperlink(paragraph, url: str, text: str | None = None, *, size=9.0):
    """Add a clickable external link, if a URL is ever needed."""
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E5C8A")
    r_pr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text or url
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, instruction: str) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instruction)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


# --------------------------------------------------------------------------- #
# Document-level helpers
# --------------------------------------------------------------------------- #
def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.4)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 13.2, NAVY, 12, 4),
        ("Heading 2", 11.3, BLUE, 8, 3),
        ("Heading 3", 10.5, BLUE, 6, 2),
    ):
        style = doc.styles[name]
        style.font.name = HEAD_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def setup_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.paragraph_format.tab_stops.clear_all()
    usable = section.page_width - section.left_margin - section.right_margin
    p.paragraph_format.tab_stops.add_tab_stop(usable, alignment=WD_TAB_ALIGNMENT.RIGHT)
    add_run(p, "SYNTHETIC DEMO DATA", bold=True, size=7.4, color=RED, font=HEAD_FONT)
    add_run(p, "\tPowerShield-FL — Conference Paper Draft", size=7.4, color=LIGHT_GRAY, font=HEAD_FONT)
    add_bottom_border(p, color="D9E2F3", size="6")

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "Page ", size=7.8, color=GRAY, font=HEAD_FONT)
    add_field(p, "PAGE")
    add_run(p, " of ", size=7.8, color=GRAY, font=HEAD_FONT)
    add_field(p, "NUMPAGES")


def add_callout(doc: Document, paragraphs: list[str], *, fill: str, border: str, width=6.86) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(width)
    cell = table.cell(0, 0)
    cell.width = Inches(width)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=110, start=150, bottom=110, end=150)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge, sz in (("top", "4"), ("bottom", "4"), ("right", "4"), ("left", "24")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), border)
        borders.append(el)
    tc_pr.append(borders)
    first = True
    for text in paragraphs:
        if not text.strip():
            continue
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.06
        add_inline(p, text, size=9.0, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    add_inline(p, text, size={1: 13.2, 2: 11.3, 3: 10.5}[level], color={1: NAVY, 2: BLUE, 3: BLUE}[level], bold=True, font=HEAD_FONT)
    if level == 1:
        add_bottom_border(p, color="D9E2F3", size="8")


def add_body(doc: Document, text: str, *, first_paragraph=False, abstract=False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    if abstract:
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.18)
        add_inline(p, text, size=9.55, color=DARK)
    else:
        add_inline(p, text, size=10.4, color=DARK)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.30)
    p.paragraph_format.first_line_indent = Inches(-0.17)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    add_run(p, "•  ", size=10.4, color=BLUE)
    add_inline(p, text, size=10.4, color=DARK)


def add_number(doc: Document, number: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.32)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    add_run(p, f"{number}.  ", size=10.4, color=BLUE)
    add_inline(p, text, size=10.4, color=DARK)


def add_equation(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    add_run(p, text, size=10.6, color=DARK, font="Cambria Math")


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    header = rows[0]
    body = rows[1:]
    n_cols = len(header)
    table = doc.add_table(rows=1, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="C9D4E3", size="4")
    set_table_cell_margins(table, top=28, start=52, bottom=28, end=52)

    # Column widths from content length; normalize to the usable width.
    usable = 6.86
    lengths = []
    for c in range(n_cols):
        vals = [header[c]] + [r[c] for r in body]
        max_len = max(len(re.sub(r"[*`]", "", v)) for v in vals)
        lengths.append(max(5, min(28, max_len)))
    weights = [max(0.32, min(2.4, 0.095 * ln)) for ln in lengths]
    total = sum(weights)
    widths = [usable * w / total for w in weights]
    widths[-1] += usable - sum(widths)

    header_cells = table.rows[0].cells
    for idx, text in enumerate(header):
        cell = header_cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F3864")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_inline(p, text, size=8.0, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True, font=HEAD_FONT)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    left_cols = {0}
    long_headers = {"Claim", "Evidence", "Status", "Result", "Parameter", "Check"}
    for idx, h in enumerate(header):
        if h in long_headers:
            left_cols.add(idx)
    numeric_headers = {
        "Precision", "Recall", "F1", "ROC-AUC", "PR-AUC", "Accuracy", "Threshold",
        "Total bytes", "Time (s)", "Train time (s)", "Inference (ms/sample)",
        "Payload (KiB)", "Parameters", "Consumers", "Assigned windows",
        "Theft windows", "Theft ratio", "Local train", "Local validation",
        "Rounds", "Clients", "Wall time (s)", "FedAvg weight", "Adaptive weight",
        "Quality", "Reliability", "Balance", "Update norm", "Rank",
        "Mean ROC-AUC drop", "Standard deviation", "Single-release ε",
        "Reduction vs none", "TP / FP / TN / FN",
    }
    center_cols = {i for i, h in enumerate(header) if h in numeric_headers} - left_cols

    for r_idx, row in enumerate(body):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for c_idx, text in enumerate(row):
            cell = cells[c_idx]
            cell.width = Inches(widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx % 2 == 1:
                set_cell_shading(cell, "F2F5F9")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx in left_cols else (WD_ALIGN_PARAGRAPH.CENTER if c_idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT)
            add_inline(p, text, size=8.1, color=DARK)

    for idx, col in enumerate(table.columns):
        col.width = Inches(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, path: Path, width_in: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))


def add_caption(doc: Document, text: str, *, kind: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind == "figure" else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2 if kind == "figure" else 7)
    p.paragraph_format.space_after = Pt(7 if kind == "figure" else 3)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = kind == "table"
    color = GRAY if kind == "figure" else NAVY
    add_inline(p, text, size=8.8 if kind == "figure" else 9.1, color=color, bold=(kind == "table"), font=HEAD_FONT)


def add_reference(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.27)
    p.paragraph_format.first_line_indent = Inches(-0.27)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(p, text, size=8.8, color=DARK)


# --------------------------------------------------------------------------- #
# Markdown parser
# --------------------------------------------------------------------------- #
def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        if re.fullmatch(r"\|?[\s:|-]+\|?", line):
            i += 1
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def parse_blockquote(lines: list[str], start: int) -> tuple[list[str], int]:
    paragraphs: list[str] = []
    current: list[str] = []
    i = start
    while i < len(lines) and lines[i].lstrip().startswith(">"):
        raw = lines[i].lstrip()[1:].strip()
        if raw:
            current.append(raw)
        elif current:
            paragraphs.append(" ".join(current))
            current = []
        i += 1
    if current:
        paragraphs.append(" ".join(current))
    return paragraphs, i


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    setup_styles(doc)
    setup_header_footer(doc)

    # Document metadata
    doc.core_properties.title = "PowerShield-FL: A Lightweight, Privacy-Aware, and Communication-Efficient Federated Learning Framework for Electricity-Theft Detection"
    doc.core_properties.subject = "Conference paper draft on synthetic electricity-theft detection"
    doc.core_properties.author = "Author Name (placeholder)"
    doc.core_properties.keywords = "federated learning; electricity theft; privacy; communication efficiency; non-IID"
    doc.core_properties.comments = "All experimental values are SYNTHETIC DEMO DATA. Full mode NOT RUN; real data DATASET NOT CONFIGURED."

    # Title / author block (everything before the first ## Abstract)
    abstract_idx = next(i for i, l in enumerate(lines) if l.strip() == "## Abstract")
    title = lines[0].lstrip("# ").strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(8)
    add_run(p, title, bold=True, size=19.0, color=NAVY, font=HEAD_FONT)
    add_bottom_border(p, color="1F3864", size="12")

    author_lines = [l.strip() for l in lines[1:abstract_idx] if l.strip() and not l.lstrip().startswith(">")]
    for idx, line in enumerate(author_lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.0
        is_name = line.startswith("**") and line.endswith("**")
        is_date = "Manuscript draft" in line
        if is_name:
            add_inline(p, line, size=11.2, color=DARK, bold=True, font=HEAD_FONT)
        elif is_date:
            p.paragraph_format.space_before = Pt(6)
            add_inline(p, line, size=9.0, color=GRAY, font=HEAD_FONT)
        else:
            add_inline(p, line, size=9.2, color=GRAY, font=HEAD_FONT)

    # Data-status callout (the first blockquote block)
    status_start = next(i for i, l in enumerate(lines) if l.lstrip().startswith(">"))
    status, _ = parse_blockquote(lines, status_start)
    add_callout(doc, status, fill="FFF4D6", border="C00000")

    i = abstract_idx
    in_abstract = False
    in_references = False
    in_appendix = False

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Headings
        if stripped.startswith("## "):
            heading = stripped[3:].strip()
            if heading == "References":
                doc.add_page_break()
                in_references = True
            elif heading.startswith("Appendix A"):
                in_appendix = True
            add_heading(doc, heading, 1)
            in_abstract = heading == "Abstract"
            i += 1
            continue
        if stripped.startswith("### "):
            add_heading(doc, stripped[4:].strip(), 2)
            in_abstract = False
            i += 1
            continue
        if stripped.startswith("#### "):
            add_heading(doc, stripped[5:].strip(), 3)
            in_abstract = False
            i += 1
            continue

        # Tables
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        # Block quotes (threat model and any other callout)
        if stripped.startswith(">"):
            paras, i = parse_blockquote(lines, i)
            if "Threat model" in " ".join(paras):
                add_callout(doc, paras, fill="EAF1F8", border="2F5D8A")
            else:
                add_callout(doc, paras, fill="F2F2F2", border="7F7F7F")
            continue

        # Images
        m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", stripped)
        if m:
            path = ROOT / m.group(1).replace("\\", "/")
            width = FIGURE_WIDTHS.get(path.name, 6.2)
            add_figure(doc, path, width)
            i += 1
            continue

        # Display equations
        if stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) > 4:
            add_equation(doc, stripped[2:-2].strip())
            i += 1
            continue
        if stripped == "$$":
            i += 1
            eq_lines = []
            while i < len(lines) and lines[i].strip() != "$$":
                eq_lines.append(lines[i].strip())
                i += 1
            if i < len(lines):
                i += 1
            add_equation(doc, " ".join(eq_lines))
            continue

        # Table / figure captions
        if stripped.startswith("**Table"):
            add_caption(doc, stripped, kind="table")
            i += 1
            continue
        if stripped.startswith("**Figure"):
            add_caption(doc, stripped, kind="figure")
            i += 1
            continue

        # References
        if in_references and re.match(r"^\[\d+\]", stripped):
            add_reference(doc, stripped)
            i += 1
            continue

        # Bullets
        if stripped.startswith("- "):
            add_bullet(doc, stripped[2:].strip())
            i += 1
            continue

        # Numbered list
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            add_number(doc, m.group(1), m.group(2).strip())
            i += 1
            continue

        # Normal paragraph
        add_body(doc, stripped, abstract=in_abstract)
        i += 1

    # Add a final note in the document properties / a visible footer already
    # covers data status.  Force field update on open.
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    doc.save(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    build()
